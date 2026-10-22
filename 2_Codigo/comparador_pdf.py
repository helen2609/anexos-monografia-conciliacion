"""
Comparador planilla de conciliacion vs PDFs del sistema del operador.

Lee los materiales instalados desde los PDF de orden de trabajo (texto nativo,
sin OCR), los compara con la fila de la planilla y arma un informe (Excel + texto)
con lo que el operador podria observar, para corregir o preparar respaldo ANTES de enviar.

Uso por consola:
    python comparador_pdf.py planilla.xlsx carpeta_con_pdfs
"""
import os
import re
import shutil
import subprocess
import sys
import zipfile
import importlib.util

import pandas as pd

# === PARSER ===
# Nombre del material en el PDF -> nombre de columna en la planilla
MAPA_MATERIALES = {
    "SOPORTE MEDIO TRAMO": "SOPORTE_MEDIO_TRAMO",
    "CAIDA TENSORIAL": "TENSORES",
    "TENSOR DROP": "TENSORES",
    "CINTA ACERO": "CINTA_ACERO",
    "ANILLAS GUIA": "ANILLAS_GUIA",
    "CAIDA DE CABLE": "CABLE_DROP",
    "CABLE DROP": "CABLE_DROP",
    "CABLE DUCTO": "CABLE_DUCTO",
    "TERMINAL SC": "TERMINAL_SC",
    "ROSETA OPTICA": "ROSETA_OPTICA",
    "CABLE UTP": "CABLE_UTP",
    "CONECTOR RJ45": "CONECTOR_RJ45",
    "CONECTOR RJ-45": "CONECTOR_RJ45",
    "PACH CORD SC-SC": "PATCH_CORD_SCSC",
    "CORDON PACH SC-SC": "PATCH_CORD_SCSC",
    "ADAPTADOR AC/DC": "ADAPTADOR_ACDC",
    "PACH CORD UTP": "PATCH_CORD_UTP",
    "CABLE HDMI": "CABLE_HDMI",
    "CABLE RCA": "CABLE_RCA",
}

# Materiales que el sistema del operador NO registra en el PDF (no se pueden comparar).
# El capuchon RJ45 no aparece en ninguno de los PDFs revisados.
NO_ESTAN_EN_SISTEMA = {"CAPUCHON_RJ45"}

_LINEA_MATERIAL = re.compile(
    r"^\s*([A-Z][A-Z0-9 /\-]+?)\s+(\d+(?:\.\d+)?)\s+(\d+(?:\.\d+)?)\s+(\d+(?:\.\d+)?)\s*$"
)


def texto_pdf(path: str) -> str:
    """Texto del PDF. Usa pdftotext si existe; si no, pdfplumber."""
    try:
        r = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True)
        if r.returncode == 0:
            return r.stdout
    except FileNotFoundError:
        pass
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        return "\n".join((p.extract_text() or "") for p in pdf.pages)


def _normalizar(nombre: str) -> str:
    return re.sub(r"\s+", " ", nombre).strip()


def leer_orden_pdf(path: str):
    """Devuelve {'solicitud', 'materiales', 'bloques_iptv'} o None si el PDF no tiene
    texto (escaneado). Las ordenes con IPTV traen un bloque de materiales por cada STB
    (seccion M); 'materiales' es la suma de la seccion L y de todos los bloques."""
    t = texto_pdf(path)
    if len(t.strip()) < 200:
        return None
    sol = re.search(r"Request\s*Number:?\s*(\d{9})", t, re.S) or re.search(r"\b(3\d{8})\b", t)

    ini, fin = t.find("L. Installed materials"), t.find("N. Field Engineer")
    zona = t[ini:fin] if 0 <= ini < fin else t
    pos_m = zona.find("M. IPTV materials installed")
    parte_l = zona if pos_m < 0 else zona[:pos_m]
    parte_m = "" if pos_m < 0 else zona[pos_m:]

    materiales = {}
    for linea in parte_l.splitlines():
        m = _LINEA_MATERIAL.match(linea)
        if m and _normalizar(m.group(1)) in MAPA_MATERIALES:
            col = MAPA_MATERIALES[_normalizar(m.group(1))]
            materiales[col] = materiales.get(col, 0) + float(m.group(3))

    # Bloques IPTV: un bloque nuevo empieza cuando se repite un nombre de material
    bloques, actual, vistos = [], {}, set()
    for linea in parte_m.splitlines():
        m = _LINEA_MATERIAL.match(linea)
        if not m:
            continue
        nombre = _normalizar(m.group(1))
        if nombre in vistos:
            bloques.append(actual)
            actual, vistos = {}, set()
        vistos.add(nombre)
        if nombre in MAPA_MATERIALES:
            col = MAPA_MATERIALES[nombre]
            actual[col] = actual.get(col, 0) + float(m.group(3))
    if vistos:
        bloques.append(actual)
    for b in bloques:
        for col, v in b.items():
            materiales[col] = materiales.get(col, 0) + v

    return {"solicitud": int(sol.group(1)) if sol else None,
            "materiales": materiales, "bloques_iptv": bloques}


# === COMPARACION ===
def describir_diferencia(p: float, s: float):
    """Clasifica una diferencia planilla (p) vs sistema (s). Devuelve (tipo, detalle)."""
    if p != int(p) and abs(p - s) < 1:
        return "REDONDEO", "La planilla tiene decimales y el sistema guarda numeros enteros"
    if p == 0:
        return "FALTA EN PLANILLA", f"El sistema registra {s:g} y la planilla 0"
    if s == 0:
        return "SOBRA EN PLANILLA", f"La planilla tiene {p:g} y el sistema no lo registra"
    if abs(p - s) == 1:
        return "DIFERENCIA DE 1", "Posible error de tipeo o de conteo"
    pe, se = str(int(p)), str(int(s))
    if len(pe) == len(se) and sum(a != b for a, b in zip(pe, se)) == 1:
        return "UN DIGITO DISTINTO", "Posible error de tipeo"
    if abs(len(pe) - len(se)) == 1:
        largo, corto = (pe, se) if len(pe) > len(se) else (se, pe)
        if any(largo[:i] + largo[i + 1:] == corto for i in range(len(largo))):
            return "DIGITO DE MAS O DE MENOS", "Posible error de tipeo"
    return "CANTIDAD DISTINTA", "Las cantidades no coinciden"


def comparar_orden(fila_planilla, orden: dict, materiales: list, precios: dict, materiales_operador: set):
    """Lista de hallazgos (dicts) de una orden: planilla vs PDF del sistema."""
    difs = []
    for mat in materiales:
        if mat in NO_ESTAN_EN_SISTEMA:
            continue
        p, s = float(fila_planilla[mat]), float(orden["materiales"].get(mat, 0))
        if abs(p - s) > 1e-9:
            difs.append((mat, p, s))

    # Materiales intercambiados (ej. ducto y drop cargados al reves en el sistema)
    cambiados = {}
    for i in range(len(difs)):
        for j in range(i + 1, len(difs)):
            a, b = difs[i], difs[j]
            if a[1] == b[2] and b[1] == a[2]:
                cambiados[a[0]], cambiados[b[0]] = b[0], a[0]

    hallazgos = []
    for mat, p, s in difs:
        tipo, detalle = describir_diferencia(p, s)
        if mat in cambiados:
            tipo = "POSIBLE CAMBIO DE MATERIAL"
            detalle = f"Los valores de {mat} y {cambiados[mat]} parecen intercambiados en el sistema"
        precio = precios.get(mat, 0.0)
        bs = round((p - s) * precio, 2)
        if precio == 0:
            efecto = "Sin efecto en el monto (material del operador)"
        elif bs > 0:
            efecto = f"Facturas {bs:.2f} Bs de mas (riesgo de multa)"
        elif bs < 0:
            efecto = f"Dejas de facturar {-bs:.2f} Bs"
        else:
            efecto = "Sin efecto en el monto"
        gravedad = "BAJA" if tipo == "REDONDEO" else ("MEDIA" if mat in materiales_operador else "ALTA")
        hallazgos.append({
            "solicitud": orden["solicitud"], "material": mat, "planilla": p, "sistema": s,
            "tipo": tipo, "detalle": detalle, "bs": bs, "efecto": efecto, "gravedad": gravedad,
        })
    return hallazgos


def alertas_iptv(fila_planilla, orden: dict):
    """Avisos preventivos sobre los bloques IPTV del sistema (uno por STB)."""
    avisos = []
    bloques = orden.get("bloques_iptv", [])
    n, esperado = len(bloques), int(fila_planilla["cantidad_iptv"] or 0)
    if n != esperado:
        avisos.append(("ALTA", f"El sistema tiene {n} bloque(s) de materiales IPTV y la planilla declara "
                               f"{esperado} IPTV. Revisar el cierre de la orden."))
    if n >= 2:
        partes = []
        for col, nombre in (("CABLE_UTP", "UTP"), ("CONECTOR_RJ45", "RJ45"), ("CABLE_HDMI", "HDMI")):
            vals = [b.get(col, 0) for b in bloques]
            if any(vals):
                partes.append(f"{nombre}: " + " + ".join(f"{v:g}" for v in vals) + f" = {sum(vals):g}")
        avisos.append(("MEDIA", f"PREVENTIVA: el sistema tiene {n} bloques IPTV (uno por STB). Si el operador compara "
                                f"contra un solo bloque, vera cifras menores a la suma. Totales en el sistema -> "
                                + "; ".join(partes) + ". Tener a mano el PDF para defender la suma."))
    return avisos


# === PROCESAR ===
def preparar_pdfs(archivos_subidos, carpeta="pdfs_sistema"):
    """Junta en una carpeta los PDFs subidos, descomprimiendo los .zip."""
    if os.path.isdir(carpeta):
        shutil.rmtree(carpeta)
    os.makedirs(carpeta)
    for nombre in archivos_subidos:
        if nombre.lower().endswith(".zip"):
            with zipfile.ZipFile(nombre) as z:
                z.extractall(carpeta)
        elif nombre.lower().endswith(".pdf"):
            shutil.copy(nombre, carpeta)
    return carpeta


def procesar_pdfs(carpeta, df, materiales, precios, materiales_operador):
    """Lee todos los PDFs de la carpeta y los compara con la planilla.
    Devuelve (hallazgos_sistema, alertas_iptv_df, info_pdf)."""
    df = df.copy()
    df["solicitud"] = df["solicitud"].astype(int)
    por_solicitud = df.drop_duplicates("solicitud").set_index("solicitud", drop=False)

    rutas = []
    for raiz, _, nombres in os.walk(carpeta):
        rutas += [os.path.join(raiz, n) for n in nombres if n.lower().endswith(".pdf")]

    hallazgos, avisos = [], []
    leidos, no_leidos, fuera = set(), [], []
    for ruta in sorted(rutas):
        orden = leer_orden_pdf(ruta)
        if orden is None or orden["solicitud"] is None:
            no_leidos.append(os.path.basename(ruta))
            continue
        sol = orden["solicitud"]
        if sol not in por_solicitud.index:
            fuera.append(sol)
            continue
        leidos.add(sol)
        fila = por_solicitud.loc[sol]
        hallazgos += comparar_orden(fila, orden, materiales, precios, materiales_operador)
        for gravedad, texto in alertas_iptv(fila, orden):
            avisos.append({"solicitud": sol, "gravedad": gravedad, "descripcion": texto})

    info = {"pdfs_en_carpeta": len(rutas), "ordenes_leidas": len(leidos), "no_leidos": no_leidos,
            "fuera_de_planilla": fuera, "ordenes_planilla": len(por_solicitud),
            "ordenes_sin_pdf": len(por_solicitud) - len(leidos), "solicitudes_leidas": sorted(leidos)}
    return pd.DataFrame(hallazgos), pd.DataFrame(avisos), info


# === INFORME ===
_ORDEN_GRAVEDAD = {"ALTA": 0, "MEDIA": 1, "BAJA": 2}
_ACCION = {
    "REDONDEO": "Poner el valor entero en la planilla (el sistema guarda enteros).",
    "POSIBLE CAMBIO DE MATERIAL": "Verificar con la orden fisica. Si el error esta en el sistema, preparar la "
                                  "justificacion para el operador (el sistema no se puede modificar).",
}
_ACCION_DEFECTO = ("Verificar con la orden fisica. Si el error es de la planilla, corregirla; si es del "
                   "sistema, preparar la justificacion para el operador.")


def construir_hallazgos(anomalias, hallazgos_sistema, alertas_iptv_df):
    """Tabla unica con todo lo que se podria observar, una fila por hallazgo."""
    filas = []
    a = anomalias[anomalias["nivel_riesgo"] != "BAJO"]
    for _, r in a.iterrows():
        filas.append({"solicitud": r["solicitud"], "origen": "Reglas sobre la planilla",
                      "gravedad": "ALTA" if r["nivel_riesgo"] == "ALTO" else "MEDIA", "material": "",
                      "planilla": "", "sistema": "", "tipo": "", "bs": "",
                      "descripcion": r["detalle_alertas"],
                      "accion_sugerida": "Revisar contra la orden fisica antes de enviar."})
    if "nota_fttb" in anomalias.columns:
        for _, r in anomalias[anomalias["nota_fttb"] != ""].iterrows():
            filas.append({"solicitud": r["solicitud"], "origen": "Posible FTTB (riesgo reducido)",
                          "gravedad": "BAJA", "material": "", "planilla": "", "sistema": "", "tipo": "",
                          "bs": "", "descripcion": r["nota_fttb"],
                          "accion_sugerida": "Confirmar que es FTTB (sin tendido aereo). Si no lo es, revisar contra la orden fisica."})
    if len(hallazgos_sistema):
        for _, r in hallazgos_sistema.iterrows():
            filas.append({"solicitud": r["solicitud"], "origen": "Planilla vs sistema (PDF)",
                          "gravedad": r["gravedad"], "material": r["material"], "planilla": r["planilla"],
                          "sistema": r["sistema"], "tipo": r["tipo"], "bs": r["bs"],
                          "descripcion": f"{r['detalle']}. {r['efecto']}",
                          "accion_sugerida": _ACCION.get(r["tipo"], _ACCION_DEFECTO)})
    if len(alertas_iptv_df):
        for _, r in alertas_iptv_df.iterrows():
            filas.append({"solicitud": r["solicitud"], "origen": "Aviso preventivo IPTV",
                          "gravedad": r["gravedad"], "material": "", "planilla": "", "sistema": "",
                          "tipo": "", "bs": "", "descripcion": r["descripcion"],
                          "accion_sugerida": "Guardar el PDF del sistema como respaldo."})
    h = pd.DataFrame(filas, columns=["solicitud", "origen", "gravedad", "material", "planilla", "sistema",
                                     "tipo", "bs", "descripcion", "accion_sugerida"])
    if len(h):
        h["_o"] = h["gravedad"].map(_ORDEN_GRAVEDAD)
        h = h.sort_values(["_o", "solicitud"]).drop(columns="_o").reset_index(drop=True)
    return h


def resumen_general(df, anomalias, hallazgos_sistema, alertas_iptv_df, info_pdf):
    hs = hallazgos_sistema if len(hallazgos_sistema) else pd.DataFrame(
        columns=["solicitud", "tipo", "bs"])
    return {
        "Ordenes en la planilla": len(df),
        "Ordenes con riesgo ALTO (reglas sobre la planilla)": int((anomalias["nivel_riesgo"] == "ALTO").sum()),
        "Ordenes con riesgo MEDIO (reglas sobre la planilla)": int((anomalias["nivel_riesgo"] == "MEDIO").sum()),
        "PDFs del sistema leidos": info_pdf["ordenes_leidas"],
        "Ordenes de la planilla sin PDF (no comparadas)": info_pdf["ordenes_sin_pdf"],
        "PDFs escaneados o ilegibles": len(info_pdf["no_leidos"]),
        "PDFs cuya solicitud no esta en la planilla": len(info_pdf["fuera_de_planilla"]),
        "Ordenes con diferencias contra el sistema": int(hs["solicitud"].nunique()) if len(hs) else 0,
        "Diferencias contra el sistema (total)": len(hs),
        "  de ellas, redondeos (decimales)": int((hs["tipo"] == "REDONDEO").sum()) if len(hs) else 0,
        "Bs que dejarias de facturar": round(float(-hs.loc[hs["bs"] < 0, "bs"].sum()), 2) if len(hs) else 0.0,
        "Bs facturados de mas (riesgo de multa)": round(float(hs.loc[hs["bs"] > 0, "bs"].sum()), 2) if len(hs) else 0.0,
        "Avisos preventivos IPTV": len(alertas_iptv_df),
        "Ordenes con alerta bajada por posible FTTB": int((anomalias["nota_fttb"] != "").sum())
                                                       if "nota_fttb" in anomalias.columns else 0,
    }


def exportar_informe(df, anomalias, hallazgos_sistema, alertas_iptv_df, resumen_tec, info_pdf,
                     nombre_base="informe_conciliacion"):
    """Crea el Excel y el texto del informe. Devuelve (ruta_xlsx, ruta_txt, resumen, hallazgos)."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    resumen = resumen_general(df, anomalias, hallazgos_sistema, alertas_iptv_df, info_pdf)
    hallazgos = construir_hallazgos(anomalias, hallazgos_sistema, alertas_iptv_df)
    resumen_df = pd.DataFrame({"Indicador": list(resumen), "Valor": list(resumen.values())})
    con_nota = anomalias["nota_fttb"] != "" if "nota_fttb" in anomalias.columns else False
    con_alertas = anomalias[(anomalias["n_alertas"] > 0) | con_nota][
        ["solicitud", "tipo_trabajo", "tipo_servicio", "puntaje_riesgo", "nivel_riesgo", "detalle_alertas"]
        + (["nota_fttb"] if "nota_fttb" in anomalias.columns else [])]
    no_leidos = pd.DataFrame({
        "archivo_o_solicitud": list(info_pdf["no_leidos"]) + [str(s) for s in info_pdf["fuera_de_planilla"]],
        "motivo": ["PDF escaneado o sin texto"] * len(info_pdf["no_leidos"])
                  + ["La solicitud no esta en la planilla"] * len(info_pdf["fuera_de_planilla"])})
    hojas = {"Resumen": resumen_df, "Hallazgos": hallazgos,
             "Planilla vs sistema": hallazgos_sistema if len(hallazgos_sistema) else pd.DataFrame(
                 columns=["solicitud", "material", "planilla", "sistema", "tipo", "detalle", "bs", "efecto", "gravedad"]),
             "Ordenes con alertas": con_alertas, "Tecnicos": resumen_tec, "PDF no leidos": no_leidos}

    ruta_xlsx, ruta_txt = f"{nombre_base}.xlsx", f"{nombre_base}.txt"
    with pd.ExcelWriter(ruta_xlsx, engine="openpyxl") as w:
        for nombre, tabla in hojas.items():
            tabla.to_excel(w, sheet_name=nombre, index=False)
            ws = w.sheets[nombre]
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F4E78")
            ws.freeze_panes = "A2"
            if ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions
            for i, col in enumerate(tabla.columns, 1):
                largo = max([len(str(col))] + [len(str(v)) for v in tabla[col].head(200)])
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, largo + 2), 70)
            for fila in ws.iter_rows(min_row=2):
                for c in fila:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        colores = {"ALTA": "F8CBAD", "MEDIA": "FFE699", "BAJA": "E2EFDA"}
        ws = w.sheets["Hallazgos"]
        for fila in ws.iter_rows(min_row=2):
            gr = fila[2].value
            if gr in colores:
                fila[2].fill = PatternFill("solid", fgColor=colores[gr])

    lineas = ["INFORME DE CONCILIACION - REVISION ANTES DE ENVIAR AL OPERADOR", "=" * 60, ""]
    lineas += [f"{k}: {v}" for k, v in resumen.items()]
    lineas += ["", "HALLAZGOS (ordenados por gravedad)", "-" * 60]
    for _, r in hallazgos.iterrows():
        mat = f" | {r['material']}: planilla={r['planilla']:g} sistema={r['sistema']:g}" if r["material"] else ""
        lineas.append(f"[{r['gravedad']}] {r['solicitud']} | {r['origen']}{mat}")
        lineas.append(f"      {r['descripcion']}")
        lineas.append(f"      -> {r['accion_sugerida']}")
    with open(ruta_txt, "w", encoding="utf-8") as f:
        f.write("\n".join(lineas) + "\n")
    return ruta_xlsx, ruta_txt, resumen, hallazgos


# === CLI ===
if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Uso: python comparador_pdf.py planilla.xlsx carpeta_con_pdfs")
        sys.exit(1)
    ruta_script = os.path.join(os.path.dirname(os.path.abspath(__file__)), "deteccion_anomalias.py")
    spec = importlib.util.spec_from_file_location("det", ruta_script)
    det = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(det)
    df_pl = det.extraer_tabla(sys.argv[1])
    anom = det.detectar_anomalias(df_pl)
    h_sis, h_iptv, info = procesar_pdfs(sys.argv[2], df_pl, det.MATERIALES, det.PRECIOS_UNITARIOS, det.MATERIALES_OPERADOR)
    x, t, res, _ = exportar_informe(df_pl, anom, h_sis, h_iptv, det.resumen_por_tecnico(df_pl), info)
    for k, v in res.items():
        print(f"{k}: {v}")
    print(f"\nGenerados: {x} y {t}")
